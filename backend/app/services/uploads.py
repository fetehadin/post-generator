"""Resume and portfolio text (FR-20, FR-21, NFR-11, NFR-12).

A resume (PDF or DOCX) or a public portfolio page becomes plain text. Contact details are
removed BEFORE anything is saved: emails, phone numbers and street-address lines. The file
or page itself is never kept. The user reads and can edit the cleaned text; only text they
mark reviewed goes into their summary card, which is what drafts use.
"""

import io
import ipaddress
import re
import socket
import time
import unicodedata
import uuid
import zipfile
from datetime import UTC, datetime
from html.parser import HTMLParser
from urllib.parse import urljoin, urlsplit

import httpx
from sqlalchemy import select

from app.core.db import user_session
from app.models import Upload

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_PAGE_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024  # a small .docx can unzip to gigabytes
MAX_TEXT = 20_000
FETCH_TIMEOUT = 10.0  # per network read
FETCH_DEADLINE = 15.0  # the whole fetch, redirects included (NFR-12)
LONG_LINE = 120  # longer lines lose only the address, not the whole line
MAX_REDIRECTS = 3

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# "name [at] mail (dot) com", "name at gmail dot com": how pages hide emails from scrapers.
_AT = r"(?:\s*[\[(]\s*at\s*[\])]\s*|\s+at\s+)"
_DOT = r"(?:\s*[\[(]\s*dot\s*[\])]\s*|\s+dot\s+|\.)"
HIDDEN_EMAIL = re.compile(rf"[\w.+-]+{_AT}[\w-]+(?:{_DOT}[a-z]{{2,}})+", re.IGNORECASE)
# Chat links carry a phone number: wa.me/251…, tel:+251…, t.me/+251…
PHONE_LINK = re.compile(r"(?:wa\.me/|tel:|t\.me/\+)\+?\d[\d\s.\-()]{6,}\d", re.IGNORECASE)
# A run of 9–15 digits, allowing +, spaces, dots, dashes and brackets between them, and an
# optional extension ("x89", "ext. 12").
PHONE = re.compile(
    r"(?<![\w/])\+?\(?\d(?:[\s.\-()]{0,2}\d){8,14}(?:\s*(?:x|ext\.?)\s*\d{1,5})?(?![\w/])",
    re.IGNORECASE,
)
# A number, at least one capitalised name word, then a street type: "221B Baker Street".
# "Built 5 Drive integrations" has no name word, so it stays.
STREET = re.compile(
    r"\b\d+[A-Za-z]?\s+(?:[A-Z][\w.'-]*\s+){1,4}"
    r"(?:Street|St\.?|Avenue|Ave\.?|Road|Rd\.?|Boulevard|Blvd\.?|Lane|Ln\.?|Drive|Dr\.?|"
    r"Way|Court|Ct\.?|Place|Pl\.?|Square|Sq\.?)\b",
)
ADDRESS_WORDS = re.compile(
    r"\b(?:P\.?\s?O\.?\s?Box|House\s+No\.?|Woreda\s*\d|Kebele\s*\d|Sub[- ]?city|"
    r"Apt\.?\s*\d|Apartment\s+\d|Postal\s+code|Zip\s*code)\b",
    re.IGNORECASE,
)


class UploadError(Exception):
    def __init__(self, code: str, message: str, status: int = 400) -> None:
        super().__init__(message)
        self.code, self.message, self.status = code, message, status


# ---- turning files and pages into text ----


def resume_text(data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise UploadError("invalid_file", "The file is larger than 5 MB.")
    if data.startswith(b"%PDF"):
        return _pdf_text(data)
    if data.startswith(b"PK"):
        return _docx_text(data)
    raise UploadError("invalid_file", "Upload your resume as a PDF or a Word (.docx) file.")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UploadError(
                "invalid_file", "That PDF is password-protected. Save a copy without one."
            )
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except UploadError:
        raise
    except Exception as e:  # pypdf raises many kinds of errors on broken files
        raise UploadError("invalid_file", "That PDF couldn't be read. Try saving it again.") from e
    if not text.strip():
        raise UploadError(
            "no_text", "That PDF has no text we can read (it may be a scan). Try a Word file."
        )
    return text


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise UploadError("invalid_file", "That Word file is too large inside.")
        doc = Document(io.BytesIO(data))
    except UploadError:
        raise
    except Exception as e:
        raise UploadError("invalid_file", "That Word file couldn't be read.") from e
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


class _PageText(HTMLParser):
    """Visible text of a page: no scripts, styles or head; blocks become new lines."""

    SKIP = {"script", "style", "noscript", "svg", "head", "template"}
    BLOCKS = {"p", "div", "li", "br", "h1", "h2", "h3", "h4", "h5", "h6", "section", "tr"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skipping:
            self.skipping -= 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skipping:
            self.parts.append(data)


def page_text(html: str) -> str:
    parser = _PageText()
    parser.feed(html)
    return "".join(parser.parts)


def tidy(text: str) -> str:
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.splitlines()]
    out = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return out[:MAX_TEXT]


# ---- removing contact details (FR-21) ----


def strip_contacts(text: str) -> tuple[str, dict[str, int]]:
    """Remove emails, phone numbers and street-address lines. Returns the cleaned text and
    how many of each were removed, so the user can see what happened."""
    removed = {"emails": 0, "phones": 0, "addresses": 0}
    kept = []
    # NFKC turns look-alikes into plain characters first: a full-width ＠ becomes @.
    for line in unicodedata.normalize("NFKC", text).splitlines():
        if STREET.search(line) or ADDRESS_WORDS.search(line):
            removed["addresses"] += 1
            if len(line) <= LONG_LINE:
                continue  # a short line with an address is the address
            line = ADDRESS_WORDS.sub("", STREET.sub("", line))
        for pattern, kind in (
            (EMAIL, "emails"),
            (HIDDEN_EMAIL, "emails"),
            (PHONE_LINK, "phones"),
            (PHONE, "phones"),
        ):
            line, n = pattern.subn("", line)
            removed[kind] += n
        kept.append(line)
    return tidy("\n".join(kept)), {k: v for k, v in removed.items() if v}


# ---- fetching a portfolio page safely (NFR-12) ----


NAT64 = ipaddress.ip_network("64:ff9b::/96")  # can wrap a private IPv4 address


def _public_address(host: str) -> str:
    """One address to connect to. Every address the name resolves to must be public: no
    localhost, private networks, link-local (cloud metadata) or reserved ranges."""
    try:
        infos = socket.getaddrinfo(host, None, type=socket.SOCK_STREAM)
        addresses = [ipaddress.ip_address(info[4][0]) for info in infos]
    except (socket.gaierror, UnicodeError, ValueError):
        raise UploadError("invalid_url", "That address isn't a public website.") from None
    for a in addresses:
        mapped = getattr(a, "ipv4_mapped", None)
        if not a.is_global or (mapped and not mapped.is_global) or a in NAT64:
            raise UploadError("invalid_url", "That address isn't a public website.")
    if not addresses:
        raise UploadError("invalid_url", "That address isn't a public website.")
    return str(addresses[0])


def check_url(url: str) -> str:
    try:
        parts = urlsplit(url.strip())
        port = parts.port  # raises ValueError for "x.com:abc" or ":99999"
    except ValueError:
        raise UploadError(
            "invalid_url", "Enter a full web address, like https://you.dev."
        ) from None
    if parts.scheme not in ("http", "https") or not parts.hostname:
        raise UploadError("invalid_url", "Enter a full web address, like https://you.dev.")
    if parts.username or parts.password or (port not in (None, 80, 443)):
        raise UploadError("invalid_url", "Use a plain public web address.")
    _public_address(parts.hostname)
    return parts.geturl()


def _pinned(url: str) -> tuple[str, dict[str, str], dict[str, str]]:
    """Connect to the address we just checked, not to whatever DNS says a moment later
    (DNS rebinding). The Host header and TLS name stay the real host, so HTTPS still
    verifies the certificate for it."""
    parts = urlsplit(url)
    ip = _public_address(parts.hostname)
    netloc = f"[{ip}]" if ":" in ip else ip
    if parts.port:
        netloc += f":{parts.port}"
    target = parts._replace(netloc=netloc).geturl()
    host = parts.hostname + (f":{parts.port}" if parts.port else "")
    extensions = {"sni_hostname": parts.hostname} if parts.scheme == "https" else {}
    return target, {"Host": host, "User-Agent": "Voxlace/1.0"}, extensions


def fetch_page(url: str, client: httpx.Client | None = None) -> str:
    """GET a public page, re-checking every redirect, with a size cap, a per-read timeout
    and a deadline for the whole fetch."""
    own = client is None
    # trust_env=False: a proxy from the environment would skip our address check.
    client = client or httpx.Client(timeout=FETCH_TIMEOUT, follow_redirects=False, trust_env=False)
    deadline = time.monotonic() + FETCH_DEADLINE
    try:
        current = check_url(url)
        for _ in range(MAX_REDIRECTS + 1):
            target, headers, extensions = _pinned(current)
            request = client.build_request("GET", target, headers=headers, extensions=extensions)
            resp = client.send(request, stream=True)
            try:
                if resp.is_redirect:
                    current = check_url(urljoin(current, resp.headers.get("location", "")))
                    continue
                if resp.status_code != 200:
                    raise UploadError("fetch_failed", f"The page answered {resp.status_code}.")
                kind = resp.headers.get("content-type", "")
                if "html" not in kind and "text/plain" not in kind:
                    raise UploadError("fetch_failed", "That address isn't a web page.")
                body = bytearray()
                for chunk in resp.iter_bytes():
                    body += chunk
                    if len(body) > MAX_PAGE_BYTES:
                        raise UploadError("fetch_failed", "That page is too large to read.")
                    if time.monotonic() > deadline:
                        raise UploadError("fetch_failed", "That page took too long to load.")
                html = bytes(body).decode(resp.encoding or "utf-8", errors="replace")
                return page_text(html) if "html" in kind else html
            finally:
                resp.close()
        raise UploadError("fetch_failed", "That address redirects too many times.")
    except httpx.HTTPError as e:
        raise UploadError(
            "fetch_failed", "That page couldn't be reached. Check the address."
        ) from e
    finally:
        if own:
            client.close()


# ---- saving and reviewing ----


def save(user_id: uuid.UUID, kind: str, source: str, raw_text: str) -> Upload:
    cleaned, removed = strip_contacts(raw_text)
    if not cleaned:
        raise UploadError("no_text", "We couldn't find any text there.")
    with user_session(user_id) as s:
        row = Upload(
            user_id=user_id, kind=kind, source=source[:500], cleaned_text=cleaned, removed=removed
        )
        s.add(row)
        s.flush()
        s.refresh(row)
        s.expunge(row)
        return row


def list_uploads(user_id: uuid.UUID) -> list[Upload]:
    with user_session(user_id) as s:
        rows = s.scalars(
            select(Upload).where(Upload.user_id == user_id).order_by(Upload.created_at.desc())
        ).all()
        s.expunge_all()
        return list(rows)


def update(
    user_id: uuid.UUID, upload_id: uuid.UUID, text: str | None, reviewed: bool | None
) -> Upload:
    """The user edits the cleaned text and/or marks it reviewed. Edited text is cleaned again,
    so a pasted phone number never slips through."""
    with user_session(user_id) as s:
        row = s.get(Upload, upload_id)
        if row is None or row.user_id != user_id:  # RLS too; rule 1
            raise UploadError("not_found", "Not found.", 404)
        if text is not None:
            cleaned, removed = strip_contacts(text)
            if not cleaned:
                raise UploadError("missing_text", "The text can't be empty. Delete it instead.")
            row.cleaned_text = cleaned
            for k, v in removed.items():
                row.removed = {**row.removed, k: row.removed.get(k, 0) + v}
        if reviewed is not None:
            row.reviewed_at = datetime.now(UTC) if reviewed else None
        s.flush()
        s.refresh(row)
        s.expunge(row)
        return row


def delete(user_id: uuid.UUID, upload_id: uuid.UUID) -> None:
    with user_session(user_id) as s:
        row = s.get(Upload, upload_id)
        if row is None or row.user_id != user_id:  # RLS too; rule 1
            raise UploadError("not_found", "Not found.", 404)
        s.delete(row)


def reviewed_texts(user_id: uuid.UUID) -> list[tuple[str, str]]:
    """(kind, text) of every reviewed upload, newest first: what the summary card may use."""
    return [(u.kind, u.cleaned_text) for u in list_uploads(user_id) if u.reviewed_at]
